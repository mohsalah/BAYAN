import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bayan.api import create_app
from bayan.config import Settings
from bayan.export import safe_cell
from bayan.models import Paper

PASSWORD = "correct horse battery"


class FakeConnector:
    name = "openalex"
    tier = "open"
    supported_filters = frozenset({"years", "document_types", "languages", "countries"})

    def __init__(self, papers):
        self.papers = papers

    def search(self, profile):
        yield from self.papers


def make(i, abstract="An abstract about reactors.", **kw):
    return Paper(source="openalex", id=f"W{i}", doi=f"10.1/{i}", title=kw.pop("title", f"Paper {i}"),
                 year=2020, native_type="article", category="peer_reviewed",
                 category_basis="journal", language="en", abstract=abstract, **kw)


@pytest.fixture
def env(tmp_path):
    papers = [make(1, first_author="Ada"), make(2, abstract=""),
              make(3, title="=HYPERLINK(\"http://x\")", abstract="+cmd")]
    app = create_app(Settings(database_path=tmp_path / "t.sqlite", search_max_results=50),
                     connector_factory=lambda lib: FakeConnector(papers))
    with TestClient(app) as client:
        client.post("/api/auth/register", json={"email": "ada@example.org", "password": PASSWORD})
        run = client.post("/api/search", json={"profile": {"query": "r", "year_from": 2020,
                                                            "year_to": 2020}}).json()
        yield client, app, run["id"]


def refs(*ids):
    return [{"source": "openalex", "id": f"W{i}"} for i in ids]


def test_abstract_is_returned_with_honest_status(env):
    client, _, run_id = env
    papers = {p["id"]: p for p in client.get(f"/api/runs/{run_id}").json()["papers"]}
    assert papers["W1"]["abstract"] == "An abstract about reactors."
    assert papers["W1"]["abstract_status"] == "supplied"
    assert papers["W2"]["abstract"] == "" and papers["W2"]["abstract_status"] == "not_supplied_by_source"


def test_csv_export_all_and_selected(env):
    client, _, run_id = env
    r = client.post(f"/api/runs/{run_id}/export", json={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="bayan-run-' in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][0] == "Title" and len(rows) == 4
    sel = client.post(f"/api/runs/{run_id}/export", json={"format": "csv", "papers": refs(2)})
    assert len(list(csv.reader(io.StringIO(sel.content.decode("utf-8-sig"))))) == 2


def test_export_neutralises_spreadsheet_formulas(env):
    client, _, run_id = env
    r = client.post(f"/api/runs/{run_id}/export", json={"format": "csv", "papers": refs(3)})
    text = r.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and "'+cmd" in text
    assert safe_cell("=1+1") == "'=1+1" and safe_cell("plain") == "plain" and safe_cell(5) == 5


def test_xlsx_export_has_papers_and_search_sheets(env):
    client, _, run_id = env
    r = client.post(f"/api/runs/{run_id}/export", json={"format": "xlsx"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Papers", "Search"]
    papers = list(wb["Papers"].values)
    assert papers[0][0] == "Title" and len(papers) == 4
    assert wb["Papers"]["A4"].data_type == "s" and str(wb["Papers"]["A4"].value).startswith("'=")
    info = {row[0]: row[1] for row in wb["Search"].values}
    assert info["Research area"] == "r" and info["Papers in this export"] == 3


def test_export_rejects_paper_outside_run_and_other_users(env):
    client, app, run_id = env
    assert client.post(f"/api/runs/{run_id}/export",
                       json={"papers": refs(99)}).status_code == 404
    other = TestClient(app)
    other.post("/api/auth/register", json={"email": "bob@example.org", "password": PASSWORD})
    assert other.post(f"/api/runs/{run_id}/export", json={"format": "csv"}).status_code == 404
    assert TestClient(app).post(f"/api/runs/{run_id}/export", json={}).status_code == 401


def test_gap_request_builds_prompt_and_packet(env):
    client, _, run_id = env
    r = client.post("/api/analysis/gaps", json={"papers": refs(1, 2), "gap_count": 2,
                                                "question": "reactor physics", "run_id": run_id})
    assert r.status_code == 200
    body = r.json()
    assert body["mode"] == "prompt" and "No analysis has run" in body["note"]
    assert "up to 2 candidate research gaps" in body["prompt"] and "unverified candidate" in body["prompt"]
    assert "[P1]" in body["prompt"] and "An abstract about reactors." in body["prompt"]
    assert "not supplied by the source" in body["prompt"]
    packet = body["packet"]
    assert packet["papers_selected"] == 2 and packet["papers_with_abstract"] == 1
    assert [p["evidence_level"] for p in packet["papers"]] == ["abstract", "metadata_only"]
    saved = client.get(f"/api/analysis/{body['id']}").json()
    assert saved["prompt"] == body["prompt"] and saved["request"]["gap_count"] == 2


def test_proposal_request_needs_gap_text_and_valid_count(env):
    client, _, _ = env
    base = {"papers": refs(1)}
    assert client.post("/api/analysis/proposals", json={**base, "gap_text": ""}).status_code == 422
    assert client.post("/api/analysis/proposals",
                       json={**base, "gap_text": "g", "proposal_count": 6}).status_code == 422
    r = client.post("/api/analysis/proposals", json={**base, "gap_text": "No benchmark for X",
                                                      "proposal_count": 3})
    assert r.status_code == 200
    prompt = r.json()["prompt"]
    assert "draft 3 distinct research proposals" in prompt and "No benchmark for X" in prompt
    assert "Do not invent novelty claims" in prompt


def test_gap_count_limits_and_empty_selection(env):
    client, _, _ = env
    assert client.post("/api/analysis/gaps", json={"papers": refs(1), "gap_count": 0}).status_code == 422
    assert client.post("/api/analysis/gaps", json={"papers": refs(1), "gap_count": 6}).status_code == 422
    assert client.post("/api/analysis/gaps", json={"papers": []}).status_code == 422


def test_analysis_is_private_and_limited_to_own_papers(env):
    client, app, _ = env
    made = client.post("/api/analysis/gaps", json={"papers": refs(1)}).json()
    other = TestClient(app)
    other.post("/api/auth/register", json={"email": "bob@example.org", "password": PASSWORD})
    assert other.get(f"/api/analysis/{made['id']}").status_code == 404
    assert other.post("/api/analysis/gaps", json={"papers": refs(1)}).status_code == 404
    assert client.post("/api/analysis/gaps", json={"papers": refs(42)}).status_code == 404


def test_prompt_includes_at_most_forty_papers(env):
    client, _, _ = env
    from bayan.analysis import MAX_PAPERS, build_packet
    papers = [{"source": "s", "id": str(i), "title": f"t{i}", "abstract": ""} for i in range(50)]
    packet = build_packet("research_gaps", papers, {}, 50)
    assert packet["papers_included"] == MAX_PAPERS and packet["truncated"] is True
