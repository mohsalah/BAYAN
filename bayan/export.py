"""CSV and Excel export of search results. Cells are neutralised against spreadsheet formulas."""

import csv
import io
import json
from datetime import UTC, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import language_name
from .regions import REGION_NAMES

COLUMNS: list[tuple[str, str, int]] = [
    ("title", "Title", 60), ("year", "Year", 7), ("first_author", "First author", 24),
    ("venue", "Venue", 30), ("doi", "DOI", 30), ("category_label", "Reference type", 20),
    ("native_type", "Source type", 16), ("document_kind", "Document kind", 12),
    ("origin", "Origin", 16), ("language", "Language code", 9),
    ("language_name", "Language", 16),
    ("countries", "Affiliation countries", 18), ("region_names", "Region", 26),
    ("geo_label", "Geography match", 18), ("source", "Library", 12),
    ("citations", "Citations", 10), ("access", "Open copy", 18), ("oa_url", "Open-copy link", 36),
    ("landing_url", "Source page", 36), ("keywords", "Keywords", 30),
    ("abstract_status", "Abstract", 14), ("abstract", "Abstract text", 90),
]
UNKNOWN = "unknown"
GEO_LABELS = {"matched": "matched", "unknown": "country unknown", "outside": "outside area"}
ACCESS_LABELS = {"open_copy_listed": "listed by source", "no_open_copy_listed": "none listed"}
ABSTRACT_LABELS = {"supplied": "supplied", "not_supplied_by_source": "not supplied by source"}
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: object) -> object:
    """Prefix text that a spreadsheet could read as a formula so it stays plain text."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def flatten(paper: dict, category_labels: dict[str, str]) -> dict:
    row = dict(paper)
    category = paper.get("category", "")
    row["category_label"] = category_labels.get(category, category)
    row["language_name"] = paper.get("language_name") or language_name(paper.get("language"))
    row["countries"] = "; ".join(paper.get("countries") or []) or UNKNOWN
    row["region_names"] = "; ".join(REGION_NAMES[r] for r in paper.get("regions") or []) or UNKNOWN
    row["geo_label"] = GEO_LABELS.get(paper.get("geo_status", ""), "not filtered")
    row["keywords"] = "; ".join(paper.get("keywords") or [])
    row["doi"] = paper.get("doi") or ""
    row["document_kind"] = paper.get("document_kind") or ""
    row["origin"] = "user-supplied" if paper.get("user_supplied") else "library search"
    row["access"] = ACCESS_LABELS.get(paper.get("access", ""), paper.get("access", ""))
    row["abstract_status"] = ABSTRACT_LABELS.get(
        paper.get("abstract_status", ""), paper.get("abstract_status", ""))
    return row


def _geography_text(geo: dict | None) -> str:
    if not geo or not geo.get("active"):
        return "no geography filter"
    return (f"{geo['mode']} on {geo['basis']}; regions version {geo['regions_version']}; "
            f"countries: {', '.join(geo['expanded_countries'])}; matched {geo['matched']}, "
            f"country unknown {geo['unknown']}, outside the area {geo['outside']}")


def to_csv(papers: list[dict], category_labels: dict[str, str]) -> bytes:
    """UTF-8 with a BOM so Excel opens Arabic, Chinese and accented text correctly."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([label for _, label, _ in COLUMNS])
    for paper in papers:
        row = flatten(paper, category_labels)
        writer.writerow([safe_cell(row.get(key, "")) for key, _, _ in COLUMNS])
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")


def to_xlsx(papers: list[dict], category_labels: dict[str, str], run: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Papers"
    ws.append([label for _, label, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for paper in papers:
        row = flatten(paper, category_labels)
        ws.append([safe_cell(row.get(key, "")) for key, _, _ in COLUMNS])
    for i, (key, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if key in ("title", "abstract"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Search")
    request, summary = run.get("request", {}), run.get("summary", {})
    profile = request.get("profile", {})
    rows = [
        ("Exported", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")),
        ("Run", run.get("id")),
        ("Research area", profile.get("query", "")),
        ("Years", f"{profile.get('year_from', '')}-{profile.get('year_to', '')}"),
        ("Reference types", ", ".join(profile.get("document_types", []))),
        ("Languages", ", ".join(profile.get("languages", []))),
        ("Query per language", json.dumps(
            (summary.get("language_queries") or {}).get("queries", {}), ensure_ascii=False)),
        ("Affiliation countries", ", ".join(profile.get("countries", []))),
        ("Regions", ", ".join(profile.get("regions", []))),
        ("Geography", _geography_text(summary.get("geography"))),
        ("Libraries", ", ".join(request.get("libraries", []))),
        ("Papers in this export", len(papers)),
        ("Papers found by the run", summary.get("found", "")),
        ("Filters a library could not apply", json.dumps(summary.get("unapplied_filters", {}))),
        ("Year shortfalls (not padded)", json.dumps(summary.get("shortfalls", {}))),
        ("Note", "Reference types come from source metadata; 'Peer-reviewed' is not verified "
                 "peer review. An open-copy link is not proof that a valid PDF is available."),
    ]
    for key, value in rows:
        info.append([key, safe_cell(value)])
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 100
    for cell in info["A"]:
        cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
